"""
Comprehensive presentation modification script:
1. Apply elegant borders and cohesive color palette across all slides
2. Insert COO.xlsx data on slides 3 & 4
3. Fix text overflow on slide 103
4. Fix title position on slides 122, 123, 124
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree
import openpyxl

prs = Presentation("/home/user/btPres/Proposal of LC Systems Upgrades v3.pptx")
slide_w = prs.slide_width   # 9144000 EMU
slide_h = prs.slide_height  # 5143500 EMU

# ─── Color Palette ────────────────────────────────────────────────────────────
# Primary brand color from presentation: #4AD080 (green)
# Cohesive palette for professional look:
COLOR_PRIMARY     = "4AD080"  # Green (existing brand)
COLOR_HEADER_BG   = "003D7A"  # Deep navy blue (professional)
COLOR_HEADER_TEXT = "FFFFFF"  # White text on header
COLOR_ROW_ALT     = "E8F4EB"  # Very light green tint
COLOR_BORDER      = "003D7A"  # Navy border
COLOR_TEXT        = "1A1A2E"  # Near-black text

# ─── TASK 1: Apply Elegant Borders to All Slides ──────────────────────────────
def add_border_to_slide(slide, idx):
    """Add a thin elegant border around the slide edge."""
    sp_tree = slide.shapes._spTree

    # Border parameters
    margin = 50000        # 0.055" margin from edge
    border_thickness = 19050  # 1.5pt border

    border_xml = (
        '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<p:nvSpPr>'
        f'<p:cNvPr id="{9900 + idx}" name="SlideBorder_{idx}"/>'
        '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
        '<p:nvPr/>'
        '</p:nvSpPr>'
        '<p:spPr>'
        '<a:xfrm>'
        f'<a:off x="{margin}" y="{margin}"/>'
        f'<a:ext cx="{slide_w - 2*margin}" cy="{slide_h - 2*margin}"/>'
        '</a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        '<a:noFill/>'
        f'<a:ln w="{border_thickness}" cap="sq">'
        '<a:solidFill>'
        f'<a:srgbClr val="{COLOR_BORDER}"/>'
        '</a:solidFill>'
        '</a:ln>'
        '</p:spPr>'
        '<p:txBody>'
        '<a:bodyPr/><a:lstStyle/>'
        '<a:p><a:endParaRPr/></a:p>'
        '</p:txBody>'
        '</p:sp>'
    )

    border_elem = etree.fromstring(border_xml)
    sp_tree.append(border_elem)


print("Adding borders to all slides...")
for idx, slide in enumerate(prs.slides):
    # Skip if border already exists
    has_border = any('SlideBorder' in s.name for s in slide.shapes if s.name)
    if not has_border:
        add_border_to_slide(slide, idx)
    if (idx + 1) % 20 == 0:
        print(f"  Processed {idx + 1}/{len(prs.slides)} slides")

print(f"  Borders added to all {len(prs.slides)} slides")


# ─── TASK 2: Insert COO.xlsx Data on Slides 3 & 4 ─────────────────────────────
print("\nReading COO.xlsx data...")
wb = openpyxl.load_workbook("/home/user/btPres/COO.xlsx")
ws = wb['Sheet1']

headers = None
data_rows = []
for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
    if row_idx == 1:
        headers = list(row)
    else:
        if any(cell is not None for cell in row):
            data_rows.append(list(row))

print(f"  Headers: {headers}")
print(f"  Data rows: {len(data_rows)}")


def make_cell_xml(text, is_header=False, is_alt=False, col_idx=0):
    """Create XML string for a table cell."""
    if is_header:
        bg_fill = f'<a:solidFill><a:srgbClr val="{COLOR_HEADER_BG}"/></a:solidFill>'
        txt_color = f'<a:solidFill><a:srgbClr val="{COLOR_HEADER_TEXT}"/></a:solidFill>'
        font_bold = 'b="1" '
        font_sz = 'sz="1000" '
        algn = "ctr"
    elif is_alt:
        bg_fill = f'<a:solidFill><a:srgbClr val="{COLOR_ROW_ALT}"/></a:solidFill>'
        txt_color = f'<a:solidFill><a:srgbClr val="{COLOR_TEXT}"/></a:solidFill>'
        font_bold = ''
        font_sz = 'sz="900" '
        algn = "ctr" if col_idx in [0, 4, 5] else "l"
    else:
        bg_fill = '<a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>'
        txt_color = f'<a:solidFill><a:srgbClr val="{COLOR_TEXT}"/></a:solidFill>'
        font_bold = ''
        font_sz = 'sz="900" '
        algn = "ctr" if col_idx in [0, 4, 5] else "l"

    # Sanitize text
    safe_text = str(text) if text is not None else ''
    safe_text = safe_text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
    safe_text = safe_text.replace('\xa0', ' ').strip()

    # Trim very long text in country-of-origin column
    if col_idx == 5 and len(safe_text) > 80:
        safe_text = safe_text[:77] + '...'

    # Smaller font for long upgrade-criteria cells
    if col_idx == 2 and len(safe_text) > 90:
        font_sz = 'sz="750" '

    border_line = (
        f'<a:ln w="12700">'
        f'<a:solidFill><a:srgbClr val="{COLOR_BORDER}"/></a:solidFill>'
        f'</a:ln>'
    )

    cell = (
        '<a:tc>'
        '<a:txBody>'
        '<a:bodyPr/><a:lstStyle/>'
        f'<a:p>'
        f'<a:pPr algn="{algn}"/>'
        f'<a:r>'
        f'<a:rPr lang="en-US" {font_sz}{font_bold}dirty="0">'
        '<a:latin typeface="Calibri"/>'
        f'{txt_color}'
        '</a:rPr>'
        f'<a:t>{safe_text}</a:t>'
        '</a:r>'
        '</a:p>'
        '</a:txBody>'
        '<a:tcPr marL="91440" marR="91440" marT="45720" marB="45720">'
        f'{bg_fill}'
        f'<a:lnL>{border_line}</a:lnL>'
        f'<a:lnR>{border_line}</a:lnR>'
        f'<a:lnT>{border_line}</a:lnT>'
        f'<a:lnB>{border_line}</a:lnB>'
        '</a:tcPr>'
        '</a:tc>'
    )
    return cell


def add_coo_table_to_slide(slide, rows_data, hdrs, label):
    """Insert the COO data table into a slide."""
    table_left  = 900000
    table_top   = 680000
    table_width = 8144000

    n_data = len(rows_data)
    header_h = 480000
    available_h = slide_h - table_top - 80000
    row_h = (available_h - header_h) // n_data
    total_h = header_h + n_data * row_h

    # Column proportions: S/N(5%), System(17%), Upgrade Criteria(41%),
    #                     Brand(13%), Brand Origin(12%), Country(12%)
    ratios = [0.05, 0.17, 0.41, 0.13, 0.12, 0.12]
    col_widths = [int(table_width * r) for r in ratios]
    col_widths[-1] = table_width - sum(col_widths[:-1])

    # Grid columns
    grid_xml = "".join(f'<a:gridCol w="{w}"/>' for w in col_widths)

    # Header row
    hdr_cells = "".join(make_cell_xml(h, is_header=True, col_idx=ci)
                        for ci, h in enumerate(hdrs))
    hdr_row = f'<a:tr h="{header_h}">{hdr_cells}</a:tr>'

    # Data rows
    data_rows_xml = ""
    for ri, row in enumerate(rows_data):
        alt = (ri % 2 == 1)
        cells = "".join(make_cell_xml(v, is_alt=alt, col_idx=ci)
                        for ci, v in enumerate(row))
        data_rows_xml += f'<a:tr h="{row_h}">{cells}</a:tr>'

    table_xml = (
        '<a:tbl>'
        '<a:tblPr firstRow="1" bandRow="1">'
        '<a:tableStyleId>{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}</a:tableStyleId>'
        '</a:tblPr>'
        f'<a:tblGrid>{grid_xml}</a:tblGrid>'
        f'{hdr_row}'
        f'{data_rows_xml}'
        '</a:tbl>'
    )

    gf_xml = (
        '<p:graphicFrame '
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<p:nvGraphicFramePr>'
        f'<p:cNvPr id="8001" name="COOTable_{label}"/>'
        '<p:cNvGraphicFramePr>'
        '<a:graphicFrameLocks noGrp="1"/>'
        '</p:cNvGraphicFramePr>'
        '<p:nvPr/>'
        '</p:nvGraphicFramePr>'
        '<p:xfrm>'
        f'<a:off x="{table_left}" y="{table_top}"/>'
        f'<a:ext cx="{table_width}" cy="{total_h}"/>'
        '</p:xfrm>'
        '<a:graphic>'
        '<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/table">'
        f'{table_xml}'
        '</a:graphicData>'
        '</a:graphic>'
        '</p:graphicFrame>'
    )

    gf_elem = etree.fromstring(gf_xml)
    slide.shapes._spTree.append(gf_elem)
    print(f"  Added COO table ({n_data} data rows) to slide {label}")


# Slide 3 gets rows 1-9, slide 4 gets rows 10-17
print(f"\nAdding COO data to slide 3 ({len(data_rows[:9])} rows)...")
add_coo_table_to_slide(prs.slides[2], data_rows[:9], headers, "Slide3")

print(f"Adding COO data to slide 4 ({len(data_rows[9:])} rows)...")
add_coo_table_to_slide(prs.slides[3], data_rows[9:], headers, "Slide4")


# ─── TASK 3: Fix Text on Slide 103 ────────────────────────────────────────────
print("\nFixing slide 103 text overflow...")
slide103 = prs.slides[102]
image_left = 6235925  # Image left edge (6.82")
right_margin = 100000  # 0.11" gap before image

for shape in slide103.shapes:
    if shape.name in ('Google Shape;875;p101', 'Google Shape;876;p101'):
        sp = shape._element
        spPr = sp.find(qn('p:spPr'))
        if spPr is None:
            continue
        xfrm = spPr.find(qn('a:xfrm'))
        if xfrm is None:
            continue
        off = xfrm.find(qn('a:off'))
        ext = xfrm.find(qn('a:ext'))
        if off is None or ext is None:
            continue
        left_val = int(off.get('x'))
        new_w = image_left - right_margin - left_val
        old_w = int(ext.get('cx'))
        ext.set('cx', str(new_w))
        print(f"  {shape.name}: width {old_w/914400:.3f}\" -> {new_w/914400:.3f}\"")


# ─── TASK 4: Fix Title Headers on Slides 122, 123, 124 ─────────────────────────
print("\nFixing title headers on slides 122, 123, 124...")

# Target title position (fits above content area that starts at y~858442)
T_LEFT   = 457200   # 0.50"
T_TOP    = 152400   # 0.167"
T_WIDTH  = 8229600  # 9.00"
T_HEIGHT = 558800   # 0.61"  => ends at 711200 (0.78"), well above content


def fix_title_position(slide, slide_num):
    for shape in slide.shapes:
        if 'Title' in shape.name:
            sp = shape._element
            spPr = sp.find(qn('p:spPr'))
            if spPr is None:
                spPr_xml = '<p:spPr xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>'
                spPr = etree.fromstring(spPr_xml)
                # Insert spPr after nvSpPr
                nvSpPr = sp.find(qn('p:nvSpPr'))
                if nvSpPr is not None:
                    nvSpPr.addnext(spPr)
                else:
                    sp.insert(0, spPr)

            xfrm = spPr.find(qn('a:xfrm'))
            if xfrm is None:
                xfrm_xml = (
                    '<a:xfrm xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                    f'<a:off x="{T_LEFT}" y="{T_TOP}"/>'
                    f'<a:ext cx="{T_WIDTH}" cy="{T_HEIGHT}"/>'
                    '</a:xfrm>'
                )
                xfrm = etree.fromstring(xfrm_xml)
                spPr.insert(0, xfrm)
                print(f"  Slide {slide_num}: Added title xfrm "
                      f"({T_LEFT/914400:.3f}\", {T_TOP/914400:.3f}\") "
                      f"{T_WIDTH/914400:.3f}\"x{T_HEIGHT/914400:.3f}\"")
            else:
                off = xfrm.find(qn('a:off'))
                ext = xfrm.find(qn('a:ext'))
                if off is not None:
                    old = (off.get('x'), off.get('y'))
                    off.set('x', str(T_LEFT))
                    off.set('y', str(T_TOP))
                    print(f"  Slide {slide_num}: Updated title offset {old} -> ({T_LEFT},{T_TOP})")
                if ext is not None:
                    ext.set('cx', str(T_WIDTH))
                    ext.set('cy', str(T_HEIGHT))


for slide_idx, slide_num in [(121, 122), (122, 123), (123, 124)]:
    fix_title_position(prs.slides[slide_idx], slide_num)


# ─── Save ─────────────────────────────────────────────────────────────────────
print("\nSaving modified presentation...")
prs.save("/home/user/btPres/Proposal of LC Systems Upgrades v3.pptx")
print("Done! File saved to: /home/user/btPres/Proposal of LC Systems Upgrades v3.pptx")
